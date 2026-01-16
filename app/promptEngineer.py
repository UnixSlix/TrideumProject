import os
import math
from docx import Document
from ollama import chat, ChatResponse
from Rag_Module import init_milvus_client, query_collection, build_context

responses = {}

base_dir = os.path.dirname(__file__)
templateFolder = os.path.join(base_dir, "templates_docs")

def titleGeneration(subject):
    TITLE_SYS_PROMPT = """You are a professional technical writer in charge of creating research paper titles. \
    Only respond with a single, clear, and concise title that accurately represents the given research topic. \
    Do not include explanations, alternatives, quotation marks, or additional text.
    """
    TITLE_USER_PROMPT = f"""Generate a single, well-crafted title for a research paper about: {subject}.\nRespond with only the title text."""
    title: ChatResponse = chat(
        model='llama3.2',
        messages=[
            {'role': 'system', 'content': TITLE_SYS_PROMPT},
            {'role': 'user',   'content': TITLE_USER_PROMPT}
        ]
    )
    return title['message']['content']


def referenceGeneration(subject):
    # (No longer used by reportGeneration, but kept for other uses)
    REF_SYS_PROMPT = """You are a professional academic writer specializing in generating citation entries. \
    Always return a single citation in the specified format, with no explanation or extra text.
    """
    REF_USER_PROMPT = f"""Generate an MLA-formatted citation for each scholarly source about {subject}. Respond with only the citation text."""
    references: ChatResponse = chat(
        model='llama3.2',
        messages=[
            {'role': 'system', 'content': REF_SYS_PROMPT},
            {'role': 'user',   'content': REF_USER_PROMPT}
        ]
    )
    return references['message']['content']


def fillTemplate(templatePath, outputPath, llamaResponse):
    report = Document(os.path.join(templateFolder, templatePath))

    # Replace both bracket-style and Jinja-style placeholders
    jinja_map = {
        'Title': 'title',
        'Abstract': 'abstract',
        'Introduction': 'introduction',
        'Topic Overview': 'overview',
        'Discussion': 'discussion',
        'Conclusion': 'conclusion',
        'References': 'references',
    }

    for para in report.paragraphs:
        text = para.text
        for key, value in llamaResponse.items():
            # old bracket style
            text = text.replace(f'[{key}]', value)

            # new Jinja style (with or without spaces)
            if key in jinja_map:
                ph = jinja_map[key]
                text = text.replace(f'{{{{ {ph} }}}}', value)
                text = text.replace(f'{{{{{ph}}}}}',   value)

        para.text = text

    out_dir = os.path.dirname(outputPath)
    os.makedirs(out_dir, exist_ok=True)
    report.save(outputPath + ".doc")


def reportGeneration(subject, length, outputPath, templateFilename):
    reportSections = ['Abstract', 'Introduction', 'Topic Overview', 'Discussion', 'Conclusion']

    # 1️⃣ Title
    responses['Title'] = titleGeneration(subject)

    # compute pages per section (round up) and convert to words
    WORDS_PER_PAGE = 350
    per_section_pages = max(1, math.ceil(int(length) / len(reportSections)))
    per_section_words = per_section_pages * WORDS_PER_PAGE

    # 2️⃣ Retrieve RAG context
    client = init_milvus_client()
    results = query_collection(client, "Trideum_Collection", subject)
    context = build_context(results)

    # 3️⃣ Section‐by‐section generation (no inline citations)
    SYS_PROMPT = f"""You are a professional technical writer helping to draft a formal academic report. \
Each section of the report must be concise, non-repetitive, well-structured, and free of clichés. Avoid vague or overly general claims. \
Focus on clarity, coherence, and depth of explanation. Use appropriate technical vocabulary and maintain a formal tone throughout. \
Ensure that generated reports meet the required number of pages specified by the user. Each paragraph should receive indentation at the beginning of the first sentence. \
Use only the information in the <context> tags below. Do not include explanations, alternatives, references, or additional text within the sections.\n<context>\n{context}\n</context>"""

    for section in reportSections:
        USER_PROMPT = f"""
Write the {section} of a formal academic report about {subject}, targeting approximately {per_section_words} words (~{per_section_pages} page{'s' if per_section_pages>1 else ''}).
Respond with only the body text—do not include section titles, headings, labels, or any introductory commentary.
Return only the pure content of the {section}, written in a formal and academic tone.
"""
        response: ChatResponse = chat(
            model='llama3.2',
            messages=[
                {'role': 'system', 'content': SYS_PROMPT},
                {'role': 'user',   'content': USER_PROMPT}
            ]
        )
        responses[section] = response['message']['content']

    # 4️⃣ References: one MLA entry per file actually used
    used_files = sorted({os.path.basename(item['filename']) for item in results})
    REF_SYS_PROMPT = """You are a professional academic writer specializing in generating citation entries. \
Always return one properly formatted MLA citation per source file, with no extra text."""
    REF_USER_PROMPT = f"""Generate an MLA-formatted citation for each of the following source files used in the report.
List each citation on its own line, and do not include any additional explanation or commentary:
{os.linesep.join(used_files)}
"""
    refs: ChatResponse = chat(
        model='llama3.2',
        messages=[
            {'role': 'system', 'content': REF_SYS_PROMPT},
            {'role': 'user',   'content': REF_USER_PROMPT}
        ]
    )
    responses['References'] = refs['message']['content']

    # 5️⃣ Fill in the .docx template selected by the user
    fillTemplate(templateFilename, outputPath, responses)
